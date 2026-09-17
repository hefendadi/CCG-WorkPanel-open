"""Contract-aware Excel parser that preserves source rows and raw cell semantics."""

from __future__ import annotations

from datetime import date, datetime, time
from decimal import Decimal
from pathlib import Path
from typing import Any

import openpyxl

from .domain import (
    ADAPTERS,
    GOODS_V1_ADAPTER,
    GOODS_V1_LEGACY_SHEET,
    GOODS_V1_REQUIRED_COLUMNS,
    GOODS_V1_SHEET,
    SKU_V1_ADAPTER,
    CellValue,
    ImportType,
    ParsedRow,
    ParsedWorkbook,
)


class ImportFileError(ValueError):
    """Raised when the workbook cannot satisfy an import adapter contract."""


def json_raw(value: Any) -> Any:
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return format(value, "f")
    return value


def _nonempty_sheet(sheet) -> bool:
    return any(cell.value is not None for row in sheet.iter_rows() for cell in row)


def _headers(sheet) -> list[str]:
    first = next(sheet.iter_rows(min_row=1, max_row=1), ())
    return ["" if cell.value is None else str(cell.value).strip() for cell in first]


def detect_import_type(workbook) -> ImportType:
    matches: list[ImportType] = []
    for import_type, adapter in ADAPTERS.items():
        if adapter.sheet not in workbook.sheetnames:
            continue
        headers = set(_headers(workbook[adapter.sheet]))
        if set(adapter.required_columns).issubset(headers):
            matches.append(import_type)
    if len(matches) != 1:
        names = ", ".join(item.value for item in matches) or "none"
        raise ImportFileError(f"Workbook import type is not uniquely identifiable: {names}")
    return matches[0]


def parse_workbook(path: str | Path, import_type: ImportType | str | None = None) -> ParsedWorkbook:
    source_path = Path(path)
    if source_path.suffix.lower() != ".xlsx" or not source_path.is_file():
        raise ImportFileError(f"Invalid Excel source: {source_path}")
    try:
        workbook = openpyxl.load_workbook(source_path, data_only=True, read_only=True)
    except Exception as exc:
        raise ImportFileError(f"Unable to read workbook: {source_path.name}") from exc

    selected = ImportType(import_type) if import_type else detect_import_type(workbook)
    adapter = GOODS_V1_ADAPTER if selected == ImportType.PRODUCT_SKU else ADAPTERS[selected]
    sheet_name = adapter.sheet
    # SKU Import V1 (Operational) uses a dedicated frozen template on its own
    # sheet; resolve it before the sheet/header checks so a V1 workbook is
    # parsed with the V1 column mapping (entity_type stays SKU).
    if selected == ImportType.SKU and SKU_V1_ADAPTER.sheet in workbook.sheetnames:
        adapter = SKU_V1_ADAPTER
        sheet_name = adapter.sheet
    # 商品导入 V2: frozen template sheet 商品导入; the legacy business sheet
    # 物料列表 (same 14 required columns) is accepted as an alternative.
    if selected == ImportType.PRODUCT_SKU:
        adapter = GOODS_V1_ADAPTER
        if GOODS_V1_SHEET in workbook.sheetnames:
            sheet_name = GOODS_V1_SHEET
        elif GOODS_V1_LEGACY_SHEET in workbook.sheetnames:
            sheet_name = GOODS_V1_LEGACY_SHEET
    if sheet_name not in workbook.sheetnames:
        raise ImportFileError(f"Required sheet is missing: {sheet_name}")
    sheet = workbook[sheet_name]
    headers = _headers(sheet)
    required_columns = (
        GOODS_V1_REQUIRED_COLUMNS
        if selected == ImportType.PRODUCT_SKU
        else adapter.required_columns
    )
    missing = [name for name in required_columns if name not in headers]
    if missing:
        raise ImportFileError(f"Required columns are missing: {', '.join(missing)}")
    duplicate_headers = {name for name in headers if name and headers.count(name) > 1}
    if duplicate_headers:
        raise ImportFileError(f"Duplicate columns are not allowed: {', '.join(sorted(duplicate_headers))}")

    # Column indexes cover every adapter column the workbook actually provides;
    # absent optional columns (e.g. explicit Product attributes in 商品导入)
    # stay out of the parsed row so normalization treats them as blank.
    indexes = {name: headers.index(name) for name in adapter.columns if name in headers}
    required_present = [name for name in required_columns if name in indexes]
    rows: list[ParsedRow] = []
    for row_number, cells in enumerate(sheet.iter_rows(min_row=2), start=2):
        selected_cells = {name: cells[indexes[name]] for name in required_present}
        if all(cell.value is None or (isinstance(cell.value, str) and not cell.value.strip()) for cell in selected_cells.values()):
            continue
        rows.append(
            ParsedRow(
                row_number=row_number,
                raw_values={
                    name: json_raw(cells[indexes[name]].value)
                    for name in required_present
                },
                cells={
                    adapter.columns[name]: CellValue(
                        cells[indexes[name]].value,
                        cells[indexes[name]].data_type,
                        cells[indexes[name]].number_format,
                        bool(cells[indexes[name]].is_date),
                    )
                    for name in indexes
                },
            )
        )

    ignored = [
        name
        for name in workbook.sheetnames
        if name != sheet_name and not _nonempty_sheet(workbook[name])
    ]
    return ParsedWorkbook(
        source_path, selected, sheet_name, rows, ignored, adapter=adapter
    )

