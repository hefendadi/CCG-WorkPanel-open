"""Actual Sales workbook parsing, MDM resolution, and monthly aggregation.

Actual Sales is sourced only from the ERP ``实发数量`` column. Inventory
movement fields are outside this module.
"""

from __future__ import annotations

import math
import unicodedata
import zipfile
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Iterable, Mapping, Protocol
from xml.etree import ElementTree

import openpyxl
from openpyxl.utils.datetime import from_excel
from sqlalchemy import select
from sqlalchemy.orm import Session, joinedload

from webapp.mdm.models import Customer, SKU


HEADER_ROW = 1
APPROVED_STATUS = "已审核"
REQUIRED_COLUMNS = (
    "日期",
    "客户",
    "单据状态",
    "物料编码",
    "物料名称",
    "实发数量",
    "仓库",
    "批号",
    "生产日期",
)


class ActualSalesFileError(ValueError):
    """The workbook shape cannot satisfy the Actual Sales contract."""


@dataclass(frozen=True)
class ActualSalesExcelRow:
    row_number: int
    values: Mapping[str, Any]


@dataclass(frozen=True)
class MdmCustomerMapping:
    customer_id: int
    customer_stable_id: str
    customer_code: str
    customer_name: str
    channel_id: int
    channel_stable_id: str
    channel_name: str


@dataclass(frozen=True)
class MdmSkuProductMapping:
    sku_id: int
    sku_stable_id: str
    sku_code: str
    product_id: int | None
    product_stable_id: str | None
    product_name: str | None


class ActualSalesMdmResolver(Protocol):
    def resolve_customers(
        self, customer_names: Iterable[str]
    ) -> Mapping[str, tuple[MdmCustomerMapping, ...]]: ...

    def resolve_skus(
        self, sku_codes: Iterable[str]
    ) -> Mapping[str, MdmSkuProductMapping]: ...


class DatabaseActualSalesMdmResolver:
    """Resolve only persisted MDM Customer/Channel and SKU/Product records."""

    def __init__(self, session: Session):
        self.session = session

    def resolve_customers(
        self, customer_names: Iterable[str]
    ) -> Mapping[str, tuple[MdmCustomerMapping, ...]]:
        names = sorted(set(customer_names))
        if not names:
            return {}
        customers = self.session.scalars(
            select(Customer)
            .options(joinedload(Customer.channel))
            .where(Customer.customer_name.in_(names))
        ).all()
        grouped: dict[str, list[MdmCustomerMapping]] = {}
        for customer in customers:
            key = _text(customer.customer_name)
            if key is None or customer.channel is None:
                continue
            grouped.setdefault(key, []).append(
                MdmCustomerMapping(
                    customer_id=customer.id,
                    customer_stable_id=customer.stable_id,
                    customer_code=customer.customer_code,
                    customer_name=customer.customer_name,
                    channel_id=customer.channel_id,
                    channel_stable_id=customer.channel.stable_id,
                    channel_name=customer.channel.channel_name,
                )
            )
        return {key: tuple(values) for key, values in grouped.items()}

    def resolve_skus(
        self, sku_codes: Iterable[str]
    ) -> Mapping[str, MdmSkuProductMapping]:
        codes = sorted(set(sku_codes))
        if not codes:
            return {}
        skus = self.session.scalars(
            select(SKU).options(joinedload(SKU.product)).where(SKU.sku_code.in_(codes))
        ).all()
        return {
            sku.sku_code: MdmSkuProductMapping(
                sku_id=sku.id,
                sku_stable_id=sku.stable_id,
                sku_code=sku.sku_code,
                product_id=sku.product_id,
                product_stable_id=sku.product.stable_id if sku.product else None,
                product_name=sku.product.product_name if sku.product else None,
            )
            for sku in skus
        }


@dataclass(frozen=True)
class NormalizedActualSalesRow:
    source_row_number: int
    sales_date: date
    customer_name: str
    document_status: str
    sku_code: str
    sku_name: str | None
    shipped_qty: Decimal
    warehouse_name: str
    batch_no: str
    production_date: date | None
    raw_values: Mapping[str, Any]
    customer_id: int
    customer_stable_id: str
    customer_code: str
    channel_id: int
    channel_stable_id: str
    channel_name: str
    sku_id: int
    sku_stable_id: str
    product_id: int
    product_stable_id: str
    product_name: str | None

    @property
    def actual_month(self) -> date:
        return self.sales_date.replace(day=1)


@dataclass(frozen=True)
class ImportFinding:
    code: str
    message: str
    field: str | None = None
    row_number: int | None = None
    scope: str = "ROW"


@dataclass(frozen=True)
class ActualSalesPreviewSummary:
    total_rows: int
    valid_rows: int
    blocking_error_count: int
    production_aggregation_rows: int
    missing_production_date_rows: int


@dataclass(frozen=True)
class ActualSalesPreviewResult:
    period_start: date | None
    period_end: date | None
    blocking_errors: tuple[ImportFinding, ...]
    valid_rows: tuple[NormalizedActualSalesRow, ...]
    summary: ActualSalesPreviewSummary

    @property
    def commit_allowed(self) -> bool:
        return not self.blocking_errors and bool(self.valid_rows)


@dataclass(frozen=True)
class ChannelProductMonthTotal:
    channel_id: int
    product_id: int
    actual_month: date
    shipped_qty: Decimal


@dataclass(frozen=True)
class ProductProductionMonthTotal:
    product_id: int
    production_date: date
    actual_month: date
    shipped_qty: Decimal


def _text(value: Any) -> str | None:
    if value is None:
        return None
    text = unicodedata.normalize("NFKC", str(value))
    text = " ".join(text.split()).strip()
    return text or None


def _code(value: Any) -> str | None:
    if isinstance(value, bool):
        raise ValueError("boolean is not a valid SKU code")
    if isinstance(value, float):
        if not math.isfinite(value):
            raise ValueError("non-finite SKU code")
        value = int(value) if value.is_integer() else value
    return _text(value)


def _date(value: Any, *, required: bool) -> date | None:
    if value is None or _text(value) is None:
        if required:
            raise ValueError("date is required")
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        converted = from_excel(value)
        return converted.date() if isinstance(converted, datetime) else converted
    text = _text(value)
    for pattern in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d", "%Y%m%d"):
        try:
            return datetime.strptime(text, pattern).date()
        except ValueError:
            pass
    raise ValueError(f"invalid date: {text}")


def _decimal(value: Any) -> Decimal:
    if isinstance(value, bool):
        raise ValueError("boolean is not a valid quantity")
    text = _text(value)
    if text is None:
        raise ValueError("quantity is required")
    try:
        result = Decimal(text.replace(",", ""))
    except InvalidOperation as exc:
        raise ValueError(f"invalid quantity: {text}") from exc
    if not result.is_finite():
        raise ValueError(f"invalid quantity: {text}")
    return result


def _last_value_row(source: Path, worksheet_path: str) -> int:
    last_row = 0
    try:
        with zipfile.ZipFile(source) as archive, archive.open(worksheet_path) as xml:
            for _event, element in ElementTree.iterparse(xml, events=("end",)):
                if element.tag.endswith("}row"):
                    if any(
                        child.tag.endswith(("}v", "}t")) and child.text not in (None, "")
                        for child in element.iter()
                    ):
                        last_row = int(element.attrib["r"])
                    element.clear()
    except (KeyError, OSError, ValueError, zipfile.BadZipFile, ElementTree.ParseError) as exc:
        raise ActualSalesFileError(
            f"Unable to inspect worksheet data range: {source.name}"
        ) from exc
    return max(last_row, HEADER_ROW)


def read_actual_sales_workbook(
    path: str | Path, *, sheet_name: str | None = None
) -> list[ActualSalesExcelRow]:
    source = Path(path)
    if source.suffix.lower() != ".xlsx" or not source.is_file():
        raise ActualSalesFileError(f"Invalid Actual Sales Excel source: {source}")
    try:
        workbook = openpyxl.load_workbook(source, read_only=True, data_only=True)
    except Exception as exc:
        raise ActualSalesFileError(f"Unable to read workbook: {source.name}") from exc
    if sheet_name is not None and sheet_name not in workbook.sheetnames:
        raise ActualSalesFileError(f"Required sheet is missing: {sheet_name}")
    sheet = workbook[sheet_name] if sheet_name else workbook.active
    header_cells = next(sheet.iter_rows(min_row=HEADER_ROW, max_row=HEADER_ROW), ())
    headers = ["" if cell.value is None else str(cell.value).strip() for cell in header_cells]
    duplicates = sorted({name for name in headers if name and headers.count(name) > 1})
    if duplicates:
        raise ActualSalesFileError(f"Duplicate columns are not allowed: {', '.join(duplicates)}")
    missing = [name for name in REQUIRED_COLUMNS if name not in headers]
    if missing:
        raise ActualSalesFileError(f"Required columns are missing: {', '.join(missing)}")
    indexes = {name: index for index, name in enumerate(headers) if name}
    last_value_row = _last_value_row(source, sheet._worksheet_path)
    rows: list[ActualSalesExcelRow] = []
    for row_number, cells in enumerate(
        sheet.iter_rows(min_row=HEADER_ROW + 1, max_row=last_value_row),
        HEADER_ROW + 1,
    ):
        values = {
            header: cells[index].value if index < len(cells) else None
            for header, index in indexes.items()
        }
        if all(_text(value) is None for value in values.values()):
            continue
        rows.append(ActualSalesExcelRow(row_number, values))
    return rows


def _finding(
    code: str,
    message: str,
    *,
    row: ActualSalesExcelRow | None = None,
    field: str | None = None,
) -> ImportFinding:
    return ImportFinding(
        code=code,
        message=message,
        field=field,
        row_number=row.row_number if row else None,
        scope="ROW" if row else "BATCH",
    )


def validate_actual_sales_import(
    rows: Iterable[ActualSalesExcelRow],
    *,
    mdm_resolver: ActualSalesMdmResolver,
) -> ActualSalesPreviewResult:
    source_rows = list(rows)
    errors: list[ImportFinding] = []
    valid_rows: list[NormalizedActualSalesRow] = []
    if not source_rows:
        errors.append(_finding("ACTUAL_SALES_NO_DATA_ROWS", "workbook contains no data rows"))

    customer_names = {
        value for row in source_rows if (value := _text(row.values.get("客户"))) is not None
    }
    sku_codes: set[str] = set()
    for row in source_rows:
        try:
            value = _code(row.values.get("物料编码"))
        except ValueError:
            continue
        if value is not None:
            sku_codes.add(value)
    customer_index = mdm_resolver.resolve_customers(customer_names)
    sku_index = mdm_resolver.resolve_skus(sku_codes)

    for source_row in source_rows:
        row_errors: list[ImportFinding] = []
        raw = dict(source_row.values)
        normalized: dict[str, Any] = {}
        parsers = {
            "sales_date": ("日期", lambda value: _date(value, required=True)),
            "customer_name": ("客户", _text),
            "document_status": ("单据状态", _text),
            "sku_code": ("物料编码", _code),
            "sku_name": ("物料名称", _text),
            "shipped_qty": ("实发数量", _decimal),
            "warehouse_name": ("仓库", _text),
            "batch_no": ("批号", _text),
            "production_date": ("生产日期", lambda value: _date(value, required=False)),
        }
        for field, (column, parser) in parsers.items():
            try:
                normalized[field] = parser(raw.get(column))
            except (ValueError, TypeError, OverflowError) as exc:
                code = (
                    "ACTUAL_SALES_DATE_INVALID"
                    if field == "sales_date"
                    else "ACTUAL_SALES_PRODUCTION_DATE_INVALID"
                    if field == "production_date"
                    else "ACTUAL_SALES_VALUE_INVALID"
                )
                row_errors.append(_finding(code, str(exc), row=source_row, field=field))
                normalized[field] = None

        for field in ("customer_name", "sku_code", "warehouse_name", "batch_no"):
            if normalized.get(field) is None:
                row_errors.append(
                    _finding(
                        "ACTUAL_SALES_REQUIRED_FIELD_MISSING",
                        f"required field is missing: {field}",
                        row=source_row,
                        field=field,
                    )
                )
        if normalized.get("document_status") != APPROVED_STATUS:
            row_errors.append(
                _finding(
                    "ACTUAL_SALES_STATUS_NOT_APPROVED",
                    f"document status must be {APPROVED_STATUS}",
                    row=source_row,
                    field="document_status",
                )
            )

        customer_matches = customer_index.get(normalized.get("customer_name"), ())
        if len(customer_matches) != 1:
            row_errors.append(
                _finding(
                    "ACTUAL_SALES_CUSTOMER_UNRESOLVED",
                    "ERP Customer must map to exactly one MDM Customer and channel",
                    row=source_row,
                    field="customer_name",
                )
            )
        sku_mapping = sku_index.get(normalized.get("sku_code"))
        if (
            sku_mapping is None
            or sku_mapping.product_id is None
            or sku_mapping.product_stable_id is None
        ):
            row_errors.append(
                _finding(
                    "ACTUAL_SALES_PRODUCT_UNRESOLVED",
                    "ERP SKU must map to an MDM SKU with Product",
                    row=source_row,
                    field="sku_code",
                )
            )
        if row_errors:
            errors.extend(row_errors)
            continue

        customer_mapping = customer_matches[0]
        assert sku_mapping is not None
        assert sku_mapping.product_id is not None
        assert sku_mapping.product_stable_id is not None
        valid_rows.append(
            NormalizedActualSalesRow(
                source_row_number=source_row.row_number,
                raw_values=raw,
                customer_id=customer_mapping.customer_id,
                customer_stable_id=customer_mapping.customer_stable_id,
                customer_code=customer_mapping.customer_code,
                channel_id=customer_mapping.channel_id,
                channel_stable_id=customer_mapping.channel_stable_id,
                channel_name=customer_mapping.channel_name,
                sku_id=sku_mapping.sku_id,
                sku_stable_id=sku_mapping.sku_stable_id,
                product_id=sku_mapping.product_id,
                product_stable_id=sku_mapping.product_stable_id,
                product_name=sku_mapping.product_name,
                **normalized,
            )
        )

    months = [row.actual_month for row in valid_rows]
    missing_production_dates = sum(row.production_date is None for row in valid_rows)
    return ActualSalesPreviewResult(
        period_start=min(months) if months else None,
        period_end=max(months) if months else None,
        blocking_errors=tuple(errors),
        valid_rows=tuple(valid_rows),
        summary=ActualSalesPreviewSummary(
            total_rows=len(source_rows),
            valid_rows=len(valid_rows),
            blocking_error_count=len(errors),
            production_aggregation_rows=len(valid_rows) - missing_production_dates,
            missing_production_date_rows=missing_production_dates,
        ),
    )


def preview_actual_sales_workbook(
    path: str | Path,
    *,
    mdm_session: Session,
    sheet_name: str | None = None,
) -> ActualSalesPreviewResult:
    try:
        rows = read_actual_sales_workbook(path, sheet_name=sheet_name)
    except ActualSalesFileError as exc:
        finding = ImportFinding(
            code="ACTUAL_SALES_FILE_CONTRACT_INVALID",
            message=str(exc),
            scope="BATCH",
        )
        return ActualSalesPreviewResult(
            period_start=None,
            period_end=None,
            blocking_errors=(finding,),
            valid_rows=(),
            summary=ActualSalesPreviewSummary(0, 0, 1, 0, 0),
        )
    return validate_actual_sales_import(
        rows,
        mdm_resolver=DatabaseActualSalesMdmResolver(mdm_session),
    )


def aggregate_channel_product_month(
    rows: Iterable[NormalizedActualSalesRow],
) -> tuple[ChannelProductMonthTotal, ...]:
    grouped: dict[tuple[int, int, date], Decimal] = {}
    for row in rows:
        key = (row.channel_id, row.product_id, row.actual_month)
        grouped[key] = grouped.get(key, Decimal("0")) + row.shipped_qty
    return tuple(
        ChannelProductMonthTotal(*key, shipped_qty)
        for key, shipped_qty in sorted(grouped.items())
    )


def aggregate_product_production_month(
    rows: Iterable[NormalizedActualSalesRow],
) -> tuple[ProductProductionMonthTotal, ...]:
    grouped: dict[tuple[int, date, date], Decimal] = {}
    for row in rows:
        if row.production_date is None:
            continue
        key = (row.product_id, row.production_date, row.actual_month)
        grouped[key] = grouped.get(key, Decimal("0")) + row.shipped_qty
    return tuple(
        ProductProductionMonthTotal(*key, shipped_qty)
        for key, shipped_qty in sorted(grouped.items())
    )
