"""Inventory Import Contract for monthly ERP workbooks.

This module stops at a validated preview.  It deliberately has no persistence,
API, UI, derived-fact, or production-lot lifecycle behavior.
"""

from __future__ import annotations

import calendar
import enum
import math
import unicodedata
import zipfile
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Iterable, Mapping, Protocol
from xml.etree import ElementTree

import openpyxl
from openpyxl.utils.datetime import from_excel
from sqlalchemy import select
from sqlalchemy.orm import Session, joinedload

from webapp.mdm.models import MasterStatus, SKU


HEADER_ROW = 4
REQUIRED_COLUMNS = (
    "物料编码",
    "物料名称",
    "仓库名称",
    "生产日期",
    "有效期至",
    "批号",
    "库存单位",
    "(期初)数量（库存）",
    "(收入)数量（库存）",
    "(发出)数量（库存）",
    "(结存)数量（库存）",
)

FIELD_BY_COLUMN = {
    "物料编码": "sku_code",
    "物料名称": "sku_name",
    "仓库名称": "warehouse_name",
    "生产日期": "production_date",
    "有效期至": "expiry_date",
    "批号": "erp_batch_no",
    "库存单位": "inventory_unit",
    "(期初)数量（库存）": "opening_qty",
    "(收入)数量（库存）": "receipt_qty",
    "(发出)数量（库存）": "issue_qty",
    "(结存)数量（库存）": "ending_qty",
}

QUANTITY_FIELDS = ("opening_qty", "receipt_qty", "issue_qty", "ending_qty")


class InventoryFileError(ValueError):
    """The workbook shape cannot satisfy the inventory import contract."""


class WarehousePlanningStatus(str, enum.Enum):
    PLANNING_AVAILABLE = "PLANNING_AVAILABLE"
    IN_TRANSIT = "IN_TRANSIT"
    NON_SELLABLE = "NON_SELLABLE"


@dataclass(frozen=True)
class WarehousePlanningPolicy:
    version: str
    warehouses: Mapping[str, WarehousePlanningStatus]

    def status_for(self, warehouse_name: str) -> WarehousePlanningStatus | None:
        normalized = _text(warehouse_name)
        for policy_name, status in self.warehouses.items():
            if _text(policy_name) == normalized:
                return status
        return None


@dataclass(frozen=True)
class InventoryExcelRow:
    row_number: int
    values: Mapping[str, Any]


@dataclass(frozen=True)
class MdmSkuMapping:
    """Read-only MDM snapshot needed by validation; no MDM logic is modified."""

    sku_id: int
    sku_stable_id: str
    sku_code: str
    sku_status: MasterStatus
    product_id: int | None
    product_stable_id: str | None
    product_status: MasterStatus | None


class InventoryMdmResolver(Protocol):
    def resolve_skus(self, sku_codes: Iterable[str]) -> Mapping[str, MdmSkuMapping]: ...


class DatabaseInventoryMdmResolver:
    """Runtime resolver backed only by persisted MDM SKU/Product records."""

    def __init__(self, session: Session):
        self.session = session

    def resolve_skus(self, sku_codes: Iterable[str]) -> Mapping[str, MdmSkuMapping]:
        codes = sorted(set(sku_codes))
        if not codes:
            return {}
        skus = self.session.scalars(
            select(SKU).options(joinedload(SKU.product)).where(SKU.sku_code.in_(codes))
        ).all()
        return {
            sku.sku_code: MdmSkuMapping(
                sku_id=sku.id,
                sku_stable_id=sku.stable_id,
                sku_code=sku.sku_code,
                sku_status=sku.status,
                product_id=sku.product_id,
                product_stable_id=sku.product.stable_id if sku.product else None,
                product_status=sku.product.status if sku.product else None,
            )
            for sku in skus
        }


@dataclass(frozen=True)
class NormalizedInventoryRow:
    source_row_number: int
    sku_code: str
    sku_name: str
    warehouse_name: str
    production_date: date
    expiry_date: date
    erp_batch_no: str
    inventory_unit: str
    opening_qty: Decimal
    receipt_qty: Decimal
    issue_qty: Decimal
    ending_qty: Decimal
    raw_values: Mapping[str, Any]
    sku_id: int
    sku_stable_id: str
    product_id: int
    product_stable_id: str
    warehouse_planning_status: WarehousePlanningStatus
    planning_policy_version: str

    @property
    def eligible_for_derived_facts(self) -> bool:
        return self.warehouse_planning_status == WarehousePlanningStatus.PLANNING_AVAILABLE

    @property
    def eligible_for_lifecycle(self) -> bool:
        return self.eligible_for_derived_facts


@dataclass(frozen=True)
class ImportFinding:
    code: str
    message: str
    field: str | None = None
    row_number: int | None = None
    scope: str = "ROW"


@dataclass(frozen=True)
class PreviewSummary:
    total_rows: int
    valid_rows: int
    blocking_error_count: int
    warning_count: int
    planning_available_rows: int
    excluded_warehouse_rows: int
    derived_fact_candidate_rows: int
    lifecycle_candidate_rows: int


@dataclass(frozen=True)
class InventoryPreviewResult:
    period_start: date
    period_end: date
    blocking_errors: tuple[ImportFinding, ...]
    warnings: tuple[ImportFinding, ...]
    valid_rows: tuple[NormalizedInventoryRow, ...]
    summary: PreviewSummary

    @property
    def commit_allowed(self) -> bool:
        """One blocking error makes the whole batch ineligible for commit."""

        return not self.blocking_errors


def _text(value: Any) -> str | None:
    if value is None:
        return None
    text = unicodedata.normalize("NFKC", str(value))
    text = " ".join(text.split()).strip()
    return text or None


def _code(value: Any) -> str | None:
    if isinstance(value, bool):
        raise ValueError("boolean is not a valid SKU code")
    if isinstance(value, float):
        if not math.isfinite(value):
            raise ValueError("non-finite SKU code")
        value = int(value) if value.is_integer() else value
    return _text(value)


def _date(value: Any) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        converted = from_excel(value)
        return converted.date() if isinstance(converted, datetime) else converted
    text = _text(value)
    if text is None:
        raise ValueError("date is required")
    for pattern in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d", "%Y%m%d"):
        try:
            return datetime.strptime(text, pattern).date()
        except ValueError:
            pass
    raise ValueError(f"invalid date: {text}")


def _decimal(value: Any) -> Decimal:
    if isinstance(value, bool):
        raise ValueError("boolean is not a valid quantity")
    text = _text(value)
    if text is None:
        # The ERP inventory report renders zero movement/balance cells as blank.
        return Decimal("0")
    try:
        result = Decimal(text.replace(",", ""))
    except InvalidOperation as exc:
        raise ValueError(f"invalid quantity: {text}") from exc
    if not result.is_finite():
        raise ValueError(f"invalid quantity: {text}")
    return result


def _complete_natural_month(period_start: date, period_end: date) -> bool:
    last_day = calendar.monthrange(period_start.year, period_start.month)[1]
    return (
        period_start.day == 1
        and period_end == date(period_start.year, period_start.month, last_day)
    )


def read_inventory_workbook(
    path: str | Path, *, sheet_name: str | None = None
) -> list[InventoryExcelRow]:
    """Read ERP rows using the frozen fourth-row header contract."""

    source = Path(path)
    if source.suffix.lower() != ".xlsx" or not source.is_file():
        raise InventoryFileError(f"Invalid inventory Excel source: {source}")
    try:
        workbook = openpyxl.load_workbook(source, read_only=True, data_only=True)
    except Exception as exc:
        raise InventoryFileError(f"Unable to read workbook: {source.name}") from exc

    if sheet_name is not None and sheet_name not in workbook.sheetnames:
        raise InventoryFileError(f"Required sheet is missing: {sheet_name}")
    sheet = workbook[sheet_name] if sheet_name else workbook.active
    header_cells = next(sheet.iter_rows(min_row=HEADER_ROW, max_row=HEADER_ROW), ())
    headers = ["" if cell.value is None else str(cell.value).strip() for cell in header_cells]
    duplicates = sorted({name for name in headers if name and headers.count(name) > 1})
    if duplicates:
        raise InventoryFileError(f"Duplicate columns are not allowed: {', '.join(duplicates)}")
    missing = [name for name in REQUIRED_COLUMNS if name not in headers]
    if missing:
        raise InventoryFileError(f"Required columns are missing: {', '.join(missing)}")

    indexes = {name: headers.index(name) for name in REQUIRED_COLUMNS}
    last_value_row = _last_value_row(source, sheet._worksheet_path)
    rows: list[InventoryExcelRow] = []
    for row_number, cells in enumerate(
        sheet.iter_rows(min_row=HEADER_ROW + 1, max_row=last_value_row),
        HEADER_ROW + 1,
    ):
        values = {name: cells[indexes[name]].value for name in REQUIRED_COLUMNS}
        if all(_text(value) is None for value in values.values()):
            continue
        rows.append(InventoryExcelRow(row_number, values))
    return rows


def _last_value_row(source: Path, worksheet_path: str) -> int:
    """Ignore Excel dimensions inflated by formatting on otherwise empty rows."""

    last_row = 0
    try:
        with zipfile.ZipFile(source) as archive, archive.open(worksheet_path) as xml:
            for _event, element in ElementTree.iterparse(xml, events=("end",)):
                if element.tag.endswith("}row"):
                    has_value = any(
                        child.tag.endswith(("}v", "}t")) and child.text not in (None, "")
                        for child in element.iter()
                    )
                    if has_value:
                        last_row = int(element.attrib["r"])
                    element.clear()
    except (KeyError, OSError, ValueError, zipfile.BadZipFile, ElementTree.ParseError) as exc:
        raise InventoryFileError(f"Unable to inspect worksheet data range: {source.name}") from exc
    return max(last_row, HEADER_ROW)


def _finding(
    code: str,
    message: str,
    *,
    row: InventoryExcelRow | None = None,
    field: str | None = None,
) -> ImportFinding:
    return ImportFinding(
        code=code,
        message=message,
        field=field,
        row_number=row.row_number if row else None,
        scope="ROW" if row else "BATCH",
    )


def validate_inventory_import(
    rows: Iterable[InventoryExcelRow],
    *,
    period_start: date,
    period_end: date,
    mdm_resolver: InventoryMdmResolver,
    warehouse_policy: WarehousePlanningPolicy,
) -> InventoryPreviewResult:
    """Normalize, validate, map, and return a non-persistent preview."""

    source_rows = list(rows)
    errors: list[ImportFinding] = []
    warnings: list[ImportFinding] = []
    valid_rows: list[NormalizedInventoryRow] = []

    if not _complete_natural_month(period_start, period_end):
        errors.append(
            _finding(
                "INVENTORY_PERIOD_NOT_COMPLETE_MONTH",
                "period_start and period_end must cover exactly one complete natural month",
            )
        )
    policy_version = _text(warehouse_policy.version)
    if policy_version is None:
        errors.append(
            _finding(
                "INVENTORY_PLANNING_POLICY_VERSION_REQUIRED",
                "planning policy version is required for derived-fact provenance",
            )
        )
        policy_version = ""
    if not source_rows:
        errors.append(_finding("INVENTORY_NO_DATA_ROWS", "workbook contains no inventory data rows"))

    source_sku_codes: set[str] = set()
    for source_row in source_rows:
        try:
            code = _code(source_row.values.get("物料编码"))
        except ValueError:
            continue
        if code:
            source_sku_codes.add(code)
    sku_index = {
        _text(code): mapping
        for code, mapping in mdm_resolver.resolve_skus(source_sku_codes).items()
    }

    for source_row in source_rows:
        row_errors: list[ImportFinding] = []
        normalized: dict[str, Any] = {}
        raw = {column: source_row.values.get(column) for column in REQUIRED_COLUMNS}

        for column in REQUIRED_COLUMNS:
            field_name = FIELD_BY_COLUMN[column]
            value = raw[column]
            try:
                if field_name == "sku_code":
                    normalized[field_name] = _code(value)
                elif field_name in {"production_date", "expiry_date"}:
                    normalized[field_name] = _date(value)
                elif field_name in QUANTITY_FIELDS:
                    normalized[field_name] = _decimal(value)
                else:
                    normalized[field_name] = _text(value)
            except (ValueError, TypeError, OverflowError) as exc:
                row_errors.append(
                    _finding(
                        "INVENTORY_VALUE_INVALID",
                        str(exc),
                        row=source_row,
                        field=field_name,
                    )
                )
                normalized[field_name] = None

        for field_name in (
            "sku_code",
            "sku_name",
            "warehouse_name",
            "production_date",
            "expiry_date",
            "erp_batch_no",
            "inventory_unit",
            *QUANTITY_FIELDS,
        ):
            if normalized.get(field_name) is None and not any(
                finding.field == field_name for finding in row_errors
            ):
                row_errors.append(
                    _finding(
                        "INVENTORY_REQUIRED_FIELD_MISSING",
                        f"required field is missing: {field_name}",
                        row=source_row,
                        field=field_name,
                    )
                )

        for field_name in ("opening_qty", "ending_qty"):
            value = normalized.get(field_name)
            if value is not None and value < 0:
                row_errors.append(
                    _finding(
                        "INVENTORY_BALANCE_QUANTITY_NEGATIVE",
                        f"{field_name} must be non-negative",
                        row=source_row,
                        field=field_name,
                    )
                )

        quantities_present = all(normalized.get(field) is not None for field in QUANTITY_FIELDS)
        if quantities_present and (
            normalized["opening_qty"] + normalized["receipt_qty"] - normalized["issue_qty"]
            != normalized["ending_qty"]
        ):
            row_errors.append(
                _finding(
                    "INVENTORY_BALANCE_MISMATCH",
                    "Opening + Receipt - Issue must equal Ending",
                    row=source_row,
                )
            )

        mapping = sku_index.get(normalized.get("sku_code"))
        if normalized.get("sku_code") is not None:
            if mapping is None or mapping.sku_status != MasterStatus.ACTIVE:
                row_errors.append(
                    _finding(
                        "INVENTORY_ACTIVE_SKU_NOT_FOUND",
                        "SKU Code must map to an ACTIVE MDM SKU",
                        row=source_row,
                        field="sku_code",
                    )
                )
            elif (
                mapping.product_id is None
                or mapping.product_stable_id is None
                or mapping.product_status != MasterStatus.ACTIVE
            ):
                row_errors.append(
                    _finding(
                        "INVENTORY_VALID_PRODUCT_NOT_FOUND",
                        "ACTIVE SKU must map to an ACTIVE Product",
                        row=source_row,
                        field="sku_code",
                    )
                )

        errors.extend(row_errors)
        if row_errors:
            continue

        assert mapping is not None and mapping.product_id is not None
        assert mapping.product_stable_id is not None
        planning_status = warehouse_policy.status_for(normalized["warehouse_name"])
        if planning_status is None:
            errors.append(
                _finding(
                    "INVENTORY_WAREHOUSE_SCOPE_UNCLASSIFIED",
                    "warehouse must be explicitly classified by the planning policy",
                    row=source_row,
                    field="warehouse_name",
                )
            )
            continue
        if planning_status != WarehousePlanningStatus.PLANNING_AVAILABLE:
            warnings.append(
                _finding(
                    "INVENTORY_NOT_PLANNING_AVAILABLE_EXCLUDED",
                    f"{planning_status.value} row is retained in raw and excluded from derived facts",
                    row=source_row,
                    field="warehouse_name",
                )
            )
        valid_rows.append(
            NormalizedInventoryRow(
                source_row_number=source_row.row_number,
                raw_values=raw,
                sku_id=mapping.sku_id,
                sku_stable_id=mapping.sku_stable_id,
                product_id=mapping.product_id,
                product_stable_id=mapping.product_stable_id,
                warehouse_planning_status=planning_status,
                planning_policy_version=policy_version,
                **normalized,
            )
        )

    # These are semantic guardrails for downstream consumers, not row defects.
    warnings.extend(
        (
            _finding(
                "INVENTORY_GROSS_ISSUE_NOT_ACTUAL_SALES",
                "Gross Issue is an inventory movement and must not be interpreted as Actual Sales",
            ),
            _finding(
                "INVENTORY_GROSS_RECEIPT_NOT_EXTERNAL_SUPPLY",
                "Gross Receipt must not be automatically classified as new external supply",
            ),
        )
    )

    planning_count = sum(row.eligible_for_derived_facts for row in valid_rows)
    summary = PreviewSummary(
        total_rows=len(source_rows),
        valid_rows=len(valid_rows),
        blocking_error_count=len(errors),
        warning_count=len(warnings),
        planning_available_rows=planning_count,
        excluded_warehouse_rows=len(valid_rows) - planning_count,
        derived_fact_candidate_rows=planning_count,
        lifecycle_candidate_rows=planning_count,
    )
    return InventoryPreviewResult(
        period_start=period_start,
        period_end=period_end,
        blocking_errors=tuple(errors),
        warnings=tuple(warnings),
        valid_rows=tuple(valid_rows),
        summary=summary,
    )


def preview_inventory_workbook(
    path: str | Path,
    *,
    period_start: date,
    period_end: date,
    mdm_session: Session,
    warehouse_policy: WarehousePlanningPolicy,
    sheet_name: str | None = None,
) -> InventoryPreviewResult:
    """Read a real .xlsx and execute the complete preview pipeline."""

    try:
        rows = read_inventory_workbook(path, sheet_name=sheet_name)
    except InventoryFileError as exc:
        error = ImportFinding(
            code="INVENTORY_FILE_CONTRACT_INVALID",
            message=str(exc),
            scope="BATCH",
        )
        return InventoryPreviewResult(
            period_start=period_start,
            period_end=period_end,
            blocking_errors=(error,),
            warnings=(),
            valid_rows=(),
            summary=PreviewSummary(
                total_rows=0,
                valid_rows=0,
                blocking_error_count=1,
                warning_count=0,
                planning_available_rows=0,
                excluded_warehouse_rows=0,
                derived_fact_candidate_rows=0,
                lifecycle_candidate_rows=0,
            ),
        )
    return validate_inventory_import(
        rows,
        period_start=period_start,
        period_end=period_end,
        mdm_resolver=DatabaseInventoryMdmResolver(mdm_session),
        warehouse_policy=warehouse_policy,
    )
