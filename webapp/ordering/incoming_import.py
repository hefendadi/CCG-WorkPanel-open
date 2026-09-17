"""Incoming Supply workbook parsing, SKU mapping, and snapshot preview."""

from __future__ import annotations

import math
import unicodedata
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Iterable, Mapping, Protocol

import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.utils.datetime import from_excel
from sqlalchemy import select
from sqlalchemy.orm import Session, joinedload

from webapp.mdm.models import SKU


HEADER_ROW = 1
FIELD_ALIASES = {
    "sku_code": ("SKU编码",),
    "sku_name": ("商品名", "商品名(中)"),
    "order_group": ("订单",),
    "factory": ("工厂",),
    "departure_port": ("发货港",),
    "destination_port": ("目的港",),
    "order_no": ("订单号",),
    "case_pack": ("入数",),
    "order_qty": ("订单数量", "订单数"),
    "shipment_qty": ("出荷数量",),
    "shipped_total": ("已出合计",),
    "difference_qty": ("出荷差额", "差额"),
    "batch_no": ("批号",),
    "etd": ("ETD",),
    "eta": ("ETA",),
    "remark": ("备注",),
    "vessel_voyage": ("船名/航次",),
    "actual_port_arrival": ("实际到港时间", "到港时间"),
    "warehouse_entry": ("进仓时间",),
    "free_container_until": ("免箱期截至时间",),
}
REQUIRED_FIELDS = ("sku_code", "sku_name", "shipment_qty", "warehouse_entry")
INHERITED_FIELDS = ("order_qty", "shipped_total", "difference_qty")


class IncomingFileError(ValueError):
    """The workbook shape cannot satisfy the Incoming V1 contract."""


@dataclass(frozen=True)
class IncomingExcelRow:
    source_sheet_name: str
    source_row_number: int
    values: Mapping[str, Any]
    raw_values: Mapping[str, Any]


@dataclass(frozen=True)
class MdmIncomingSkuMapping:
    sku_id: int
    sku_stable_id: str
    sku_code: str
    product_id: int | None
    product_stable_id: str | None
    product_name: str | None


class IncomingMdmResolver(Protocol):
    def resolve_skus(
        self, sku_codes: Iterable[str]
    ) -> Mapping[str, MdmIncomingSkuMapping]: ...


class DatabaseIncomingMdmResolver:
    def __init__(self, session: Session):
        self.session = session

    def resolve_skus(
        self, sku_codes: Iterable[str]
    ) -> Mapping[str, MdmIncomingSkuMapping]:
        codes = sorted(set(sku_codes))
        if not codes:
            return {}
        skus = self.session.scalars(
            select(SKU).options(joinedload(SKU.product)).where(SKU.sku_code.in_(codes))
        ).all()
        return {
            sku.sku_code: MdmIncomingSkuMapping(
                sku_id=sku.id,
                sku_stable_id=sku.stable_id,
                sku_code=sku.sku_code,
                product_id=sku.product_id,
                product_stable_id=sku.product.stable_id if sku.product else None,
                product_name=sku.product.product_name if sku.product else None,
            )
            for sku in skus
        }


@dataclass(frozen=True)
class NormalizedIncomingRow:
    source_sheet_name: str
    source_row_number: int
    sku_code: str
    sku_name: str | None
    sku_id: int
    sku_stable_id: str
    product_id: int
    product_stable_id: str
    product_name: str | None
    incoming_qty: Decimal | None
    expected_arrival_date: date | None
    warehouse_entry_recorded: bool
    batch_no: str | None
    order_group: str | None
    order_no: str | None
    order_qty: Decimal | None
    shipped_total: Decimal | None
    difference_qty: Decimal | None
    raw_values: Mapping[str, Any]

    @property
    def is_current_future_supply(self) -> bool:
        return self.incoming_qty is not None and not self.warehouse_entry_recorded


@dataclass(frozen=True)
class IncomingFinding:
    code: str
    message: str
    source_sheet_name: str | None = None
    source_row_number: int | None = None
    field: str | None = None


@dataclass(frozen=True)
class ProductFutureSupplyTotal:
    product_id: int
    product_stable_id: str
    product_name: str | None
    incoming_qty: Decimal


@dataclass(frozen=True)
class IncomingPreviewSummary:
    total_rows: int
    valid_rows: int
    blocking_error_count: int
    shipment_rows: int
    current_future_supply_rows: int
    entered_shipment_rows: int
    unshipped_rows: int
    current_future_supply_qty: Decimal


@dataclass(frozen=True)
class IncomingPreviewResult:
    blocking_errors: tuple[IncomingFinding, ...]
    valid_rows: tuple[NormalizedIncomingRow, ...]
    product_totals: tuple[ProductFutureSupplyTotal, ...]
    summary: IncomingPreviewSummary

    @property
    def commit_allowed(self) -> bool:
        return not self.blocking_errors and bool(self.valid_rows)


def _text(value: Any) -> str | None:
    if value is None:
        return None
    text = unicodedata.normalize("NFKC", str(value))
    text = " ".join(text.split()).strip()
    return text or None


def _code(value: Any) -> str | None:
    if isinstance(value, bool):
        raise ValueError("boolean is not a valid SKU code")
    if isinstance(value, float):
        if not math.isfinite(value):
            raise ValueError("non-finite SKU code")
        value = int(value) if value.is_integer() else value
    return _text(value)


def _optional_decimal(value: Any) -> Decimal | None:
    text = _text(value)
    if text is None:
        return None
    if isinstance(value, bool):
        raise ValueError("boolean is not a valid quantity")
    try:
        result = Decimal(text.replace(",", ""))
    except InvalidOperation as exc:
        raise ValueError(f"invalid quantity: {text}") from exc
    if not result.is_finite():
        raise ValueError(f"invalid quantity: {text}")
    return result


def _optional_date(value: Any) -> date | None:
    text = _text(value)
    if text is None:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        converted = from_excel(value)
        return converted.date() if isinstance(converted, datetime) else converted
    for pattern in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d", "%Y%m%d"):
        try:
            return datetime.strptime(text, pattern).date()
        except ValueError:
            pass
    raise ValueError(f"invalid date: {text}")


def _header_indexes(headers: list[Any], sheet_name: str) -> dict[str, int | None]:
    normalized = [_text(value) for value in headers]
    duplicates = sorted(
        {value for value in normalized if value and normalized.count(value) > 1}
    )
    if duplicates:
        raise IncomingFileError(
            f"{sheet_name}: duplicate columns are not allowed: {', '.join(duplicates)}"
        )
    indexes: dict[str, int | None] = {}
    for field, aliases in FIELD_ALIASES.items():
        indexes[field] = next(
            (normalized.index(alias) + 1 for alias in aliases if alias in normalized),
            None,
        )
    missing = [field for field in REQUIRED_FIELDS if indexes[field] is None]
    if missing:
        labels = [FIELD_ALIASES[field][0] for field in missing]
        raise IncomingFileError(
            f"{sheet_name}: required columns are missing: {', '.join(labels)}"
        )
    return indexes


def _merged_cell_values(sheet) -> dict[tuple[int, int], Any]:
    values: dict[tuple[int, int], Any] = {}
    for merged_range in sheet.merged_cells.ranges:
        value = sheet.cell(merged_range.min_row, merged_range.min_col).value
        for row_number in range(merged_range.min_row, merged_range.max_row + 1):
            for column_number in range(merged_range.min_col, merged_range.max_col + 1):
                values[(row_number, column_number)] = value
    return values


def read_incoming_workbook(path: str | Path) -> list[IncomingExcelRow]:
    source = Path(path)
    if source.suffix.lower() != ".xlsx" or not source.is_file():
        raise IncomingFileError(f"Invalid Incoming Excel source: {source}")
    try:
        workbook = openpyxl.load_workbook(source, data_only=True)
    except Exception as exc:
        raise IncomingFileError(f"Unable to read workbook: {source.name}") from exc

    rows: list[IncomingExcelRow] = []
    candidate_sheets = 0
    for sheet in workbook.worksheets:
        headers = [cell.value for cell in sheet[HEADER_ROW]]
        normalized_headers = {_text(value) for value in headers if _text(value)}
        product_headers = set(FIELD_ALIASES["sku_name"])
        if not normalized_headers.intersection(product_headers):
            continue
        candidate_sheets += 1
        indexes = _header_indexes(headers, sheet.title)
        merged_values = _merged_cell_values(sheet)
        raw_labels = [
            _text(value) or f"__column_{get_column_letter(index)}"
            for index, value in enumerate(headers, 1)
        ]
        sheet_rows: list[IncomingExcelRow] = []
        inherited: dict[tuple[str | None, str | None, str | None], dict[str, Any]] = {}
        for row_number in range(HEADER_ROW + 1, sheet.max_row + 1):
            raw_values = {
                raw_labels[column - 1]: sheet.cell(row_number, column).value
                for column in range(1, len(headers) + 1)
            }
            values: dict[str, Any] = {}
            for field, column in indexes.items():
                values[field] = (
                    merged_values.get((row_number, column), sheet.cell(row_number, column).value)
                    if column is not None
                    else None
                )
            if all(
                _text(values.get(field)) is None
                for field in ("sku_code", "sku_name", "shipment_qty", "order_qty")
            ):
                continue
            group_key = (
                _text(values.get("order_group")),
                _text(values.get("order_no")),
                _code(values.get("sku_code")),
            )
            group_values = inherited.setdefault(group_key, {})
            for field in INHERITED_FIELDS:
                if _text(values.get(field)) is not None:
                    group_values[field] = values[field]
                elif field in group_values:
                    values[field] = group_values[field]
            sheet_rows.append(
                IncomingExcelRow(sheet.title, row_number, values, raw_values)
            )
        rows.extend(sheet_rows)
    if candidate_sheets == 0:
        raise IncomingFileError("workbook contains no Incoming data sheets")
    return rows


def _finding(
    code: str,
    message: str,
    *,
    row: IncomingExcelRow | None = None,
    field: str | None = None,
) -> IncomingFinding:
    return IncomingFinding(
        code=code,
        message=message,
        source_sheet_name=row.source_sheet_name if row else None,
        source_row_number=row.source_row_number if row else None,
        field=field,
    )


def aggregate_product_current_future_supply(
    rows: Iterable[NormalizedIncomingRow],
) -> tuple[ProductFutureSupplyTotal, ...]:
    grouped: dict[int, ProductFutureSupplyTotal] = {}
    for row in rows:
        if not row.is_current_future_supply:
            continue
        current = grouped.get(row.product_id)
        grouped[row.product_id] = ProductFutureSupplyTotal(
            product_id=row.product_id,
            product_stable_id=row.product_stable_id,
            product_name=row.product_name,
            incoming_qty=(current.incoming_qty if current else Decimal("0"))
            + row.incoming_qty,
        )
    return tuple(grouped[key] for key in sorted(grouped))


def validate_incoming_import(
    rows: Iterable[IncomingExcelRow],
    *,
    mdm_resolver: IncomingMdmResolver,
) -> IncomingPreviewResult:
    source_rows = list(rows)
    errors: list[IncomingFinding] = []
    valid_rows: list[NormalizedIncomingRow] = []
    if not source_rows:
        errors.append(_finding("INCOMING_NO_DATA_ROWS", "workbook contains no Incoming rows"))

    source_codes: set[str] = set()
    for row in source_rows:
        try:
            code = _code(row.values.get("sku_code"))
        except ValueError:
            continue
        if code:
            source_codes.add(code)
    sku_index = mdm_resolver.resolve_skus(source_codes)

    for row in source_rows:
        row_errors: list[IncomingFinding] = []
        try:
            sku_code = _code(row.values.get("sku_code"))
        except ValueError as exc:
            sku_code = None
            row_errors.append(_finding("INCOMING_SKU_INVALID", str(exc), row=row, field="sku_code"))
        if sku_code is None:
            row_errors.append(
                _finding(
                    "INCOMING_SKU_REQUIRED",
                    "SKU编码 is required for formal mapping",
                    row=row,
                    field="sku_code",
                )
            )
        mapping = sku_index.get(sku_code)
        if mapping is None or mapping.product_id is None or mapping.product_stable_id is None:
            row_errors.append(
                _finding(
                    "INCOMING_PRODUCT_UNRESOLVED",
                    "SKU编码 must map to an MDM SKU with Product",
                    row=row,
                    field="sku_code",
                )
            )

        parsed: dict[str, Any] = {}
        for field in ("shipment_qty", "order_qty", "shipped_total", "difference_qty"):
            try:
                parsed[field] = _optional_decimal(row.values.get(field))
            except ValueError as exc:
                parsed[field] = None
                row_errors.append(
                    _finding("INCOMING_QUANTITY_INVALID", str(exc), row=row, field=field)
                )
        if parsed["shipment_qty"] is not None and parsed["shipment_qty"] < 0:
            row_errors.append(
                _finding(
                    "INCOMING_QUANTITY_NEGATIVE",
                    "出荷数量 must be non-negative",
                    row=row,
                    field="shipment_qty",
                )
            )
        try:
            expected_arrival_date = _optional_date(row.values.get("eta"))
        except ValueError as exc:
            expected_arrival_date = None
            row_errors.append(_finding("INCOMING_ETA_INVALID", str(exc), row=row, field="eta"))

        if row_errors:
            errors.extend(row_errors)
            continue
        assert mapping is not None
        assert mapping.product_id is not None
        assert mapping.product_stable_id is not None
        valid_rows.append(
            NormalizedIncomingRow(
                source_sheet_name=row.source_sheet_name,
                source_row_number=row.source_row_number,
                sku_code=sku_code,
                sku_name=_text(row.values.get("sku_name")),
                sku_id=mapping.sku_id,
                sku_stable_id=mapping.sku_stable_id,
                product_id=mapping.product_id,
                product_stable_id=mapping.product_stable_id,
                product_name=mapping.product_name,
                incoming_qty=parsed["shipment_qty"],
                expected_arrival_date=expected_arrival_date,
                warehouse_entry_recorded=_text(row.values.get("warehouse_entry")) is not None,
                batch_no=_text(row.values.get("batch_no")),
                order_group=_text(row.values.get("order_group")),
                order_no=_text(row.values.get("order_no")),
                order_qty=parsed["order_qty"],
                shipped_total=parsed["shipped_total"],
                difference_qty=parsed["difference_qty"],
                raw_values=row.raw_values,
            )
        )

    product_totals = aggregate_product_current_future_supply(valid_rows)
    shipment_rows = sum(row.incoming_qty is not None for row in valid_rows)
    future_rows = sum(row.is_current_future_supply for row in valid_rows)
    entered_rows = sum(
        row.incoming_qty is not None and row.warehouse_entry_recorded for row in valid_rows
    )
    return IncomingPreviewResult(
        blocking_errors=tuple(errors),
        valid_rows=tuple(valid_rows),
        product_totals=product_totals,
        summary=IncomingPreviewSummary(
            total_rows=len(source_rows),
            valid_rows=len(valid_rows),
            blocking_error_count=len(errors),
            shipment_rows=shipment_rows,
            current_future_supply_rows=future_rows,
            entered_shipment_rows=entered_rows,
            unshipped_rows=len(valid_rows) - shipment_rows,
            current_future_supply_qty=sum(
                (total.incoming_qty for total in product_totals), Decimal("0")
            ),
        ),
    )


def preview_incoming_workbook(
    path: str | Path,
    *,
    mdm_session: Session,
) -> IncomingPreviewResult:
    try:
        rows = read_incoming_workbook(path)
    except IncomingFileError as exc:
        finding = _finding("INCOMING_FILE_CONTRACT_INVALID", str(exc))
        return IncomingPreviewResult(
            blocking_errors=(finding,),
            valid_rows=(),
            product_totals=(),
            summary=IncomingPreviewSummary(0, 0, 1, 0, 0, 0, 0, Decimal("0")),
        )
    return validate_incoming_import(
        rows,
        mdm_resolver=DatabaseIncomingMdmResolver(mdm_session),
    )
